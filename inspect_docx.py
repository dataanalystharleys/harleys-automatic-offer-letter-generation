import docx
import zipfile

doc_path = r"C:\Users\harle\Downloads\Offer Letter Harley.docx"
doc = docx.Document(doc_path)

print("=== DOCUMENT PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    print(f"P{i} [style={p.style.name}]: '{p.text}'")
    for j, r in enumerate(p.runs):
        if r.text:
            print(f"   R{j}: '{r.text}' | bold={r.bold} | italic={r.italic} | font={r.font.name} | size={r.font.size.pt if r.font.size else None} | color={r.font.color.rgb if r.font.color else None}")

print("\n=== SECTIONS ===")
for s_idx, s in enumerate(doc.sections):
    print(f"Section {s_idx}:")
    for p in s.header.paragraphs:
        if p.text.strip():
            print(f"  Header P: '{p.text}'")
    for p in s.footer.paragraphs:
        if p.text.strip():
            print(f"  Footer P: '{p.text}'")

with zipfile.ZipFile(doc_path) as z:
    for name in z.namelist():
        if name.startswith("word/media/"):
            print(f"Media: {name} ({z.getinfo(name).file_size} bytes)")

