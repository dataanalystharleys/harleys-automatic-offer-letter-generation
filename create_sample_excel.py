import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

excel_path = "sample_candidates.xlsx"

data = [
    {
        "Candidate Name": "Abhishek Rajput",
        "Salutation": "Mr.",
        "Designation": "COO",
        "Annual Salary (INR)": 5000000,
        "Salary in Words": "Fifty Lakh only/-",
        "Date of Joining": "20-10-2026",
        "Offer Date": "24-Sep-2026",
        "Address": "Marine Drive Patna, Bihar 843323",
        "Phone": "+91 9267778873",
        "Email": "abhi@gmail.com",
        "Reporting Location": "Nanakramguda",
        "Reporting Time": "10.00 AM",
        "HR Signatory": "Sruti Associate - HR"
    },
    {
        "Candidate Name": "Priya Sharma",
        "Salutation": "Ms.",
        "Designation": "Senior Full Stack Developer",
        "Annual Salary (INR)": 1800000,
        "Salary in Words": "Eighteen Lakh only/-",
        "Date of Joining": "01-11-2026",
        "Offer Date": "12-Oct-2026",
        "Address": "Plot 42, Hitech City, Hyderabad, Telangana 500081",
        "Phone": "+91 9876501234",
        "Email": "priya.sharma@example.com",
        "Reporting Location": "Nanakramguda",
        "Reporting Time": "10.00 AM",
        "HR Signatory": "Sruti Associate - HR"
    },
    {
        "Candidate Name": "Rahul Verma",
        "Salutation": "Mr.",
        "Designation": "Marketing Manager",
        "Annual Salary (INR)": 1200000,
        "Salary in Words": "Twelve Lakh only/-",
        "Date of Joining": "15-11-2026",
        "Offer Date": "12-Oct-2026",
        "Address": "Flat 301, Whitefield, Bangalore, Karnataka 560066",
        "Phone": "+91 9123456789",
        "Email": "rahul.verma@example.com",
        "Reporting Location": "Nanakramguda",
        "Reporting Time": "10.00 AM",
        "HR Signatory": "Sruti Associate - HR"
    }
]

df = pd.DataFrame(data)
df.to_excel(excel_path, index=False)

# Add styling to Excel file
wb = openpyxl.load_workbook(excel_path)
ws = wb.active

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
border_thin = Border(
    left=Side(style='thin', color='DDDDDD'),
    right=Side(style='thin', color='DDDDDD'),
    top=Side(style='thin', color='DDDDDD'),
    bottom=Side(style='thin', color='DDDDDD')
)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.font = Font(name="Calibri", size=10)
        cell.border = border_thin
        if cell.column_letter in ['D']: # Salary
            cell.number_format = '#,##,##0'

for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = col[0].column_letter
    ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

wb.save(excel_path)
print("Excel template created and styled successfully!")

