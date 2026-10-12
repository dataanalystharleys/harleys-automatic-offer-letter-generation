import os
import shutil
import docx
from docxtpl import DocxTemplate

source_file = r"C:\Users\harle\Downloads\Roshan Mokal Appointment letter.docx"
target_dir = r"C:\Users\harle\.gemini\antigravity\scratch\harleys-automatic-offer-letter-generation"
template_path = os.path.join(target_dir, "template_appointment_letter.docx")

shutil.copyfile(source_file, template_path)
doc = docx.Document(template_path)

# P3: Candidate Name
p3 = doc.paragraphs[3]
p3.runs[0].text = "{{ candidate_name }},"

# P4, P5, P6: Address
p4 = doc.paragraphs[4]
p4.runs[0].text = "{{ address }}"

p5 = doc.paragraphs[5]
p5.runs[0].text = ""

p6 = doc.paragraphs[6]
p6.runs[0].text = ""

# P7: Email
p7 = doc.paragraphs[7]
p7.runs[0].text = "Email: {{ email }}"

# P8: Mobile
p8 = doc.paragraphs[8]
p8.runs[0].text = "Mob: {{ phone }}"

# P9: Date
p9 = doc.paragraphs[9]
p9.runs[0].text = "       \t    Date"
p9.runs[1].text = ": {{ appointment_date }}"

# P12: Dear [Candidate Name]
p12 = doc.paragraphs[12]
p12.runs[0].text = "Dear {{ salutation }} {{ candidate_name }},"

# P14: Appointment intro
p14 = doc.paragraphs[14]
p14.runs[0].text = "Congratulations!  We are pleased to appoint you as "
p14.runs[1].text = "“{{ designation }}” "
p14.runs[2].text = "in "
p14.runs[3].text = "“{{ company_name }}” "
p14.runs[4].text = "w.e.f. "
p14.runs[5].text = "{{ joining_date }}"
p14.runs[6].text = ". Your work location will be at "
p14.runs[7].text = "{{ work_location }}."

# P16: CTC clause
p16 = doc.paragraphs[16]
p16.runs[0].text = "The position carries a Cost to Company (CTC) of "
p16.runs[1].text = "Rs. {{ salary_amount }} /- ({{ salary_in_words }}) per annum. "
p16.runs[2].text = "The CTC includes applicable statutory contributions towards Employees’ Provident Fund, Employees State Insurance, Professional tax and Income Tax etc.  "

# P18: Probation
p18 = doc.paragraphs[18]
p18.runs[0].text = "You will be on probation for a period of "
p18.runs[1].text = "{{ probation_period }}. "
p18.runs[2].text = "After completion of the probation period, based on the satisfactory professional performance and background verification checks, you will be offered regular employee status via confirmation letter."

# P30: Signatory Name
p30 = doc.paragraphs[30]
p30.runs[0].text = "{{ signatory_name }}"

# P31: Signatory Title
p31 = doc.paragraphs[31]
p31.runs[0].text = "{{ signatory_title }}"

# P37: Annexure A Compensation
p37 = doc.paragraphs[37]
p37.runs[0].text = "1. Compensation:"
p37.runs[1].text = " The position carries a Cost to Company (CTC) "
p37.runs[2].text = ""
p37.runs[3].text = "Rs. {{ salary_amount }} /- ({{ salary_in_words }}) per annum."
p37.runs[4].text = " "
p37.runs[5].text = "CTC includes applicable statutory contributions towards Employees’ Provident Fund, Employees State Insurance, Professional tax and Income Tax. Salary will not be paid in case of leaving the Company within a month from the date of joining. "

# P39: Annexure A Probation
p39 = doc.paragraphs[39]
p39.runs[0].text = "2. Probationary Period:"
p39.runs[1].text = " You will be on Probation for a period of "
p39.runs[2].text = "{{ probation_period }}"
p39.runs[3].text = " from the date of your joining. At the end of the probation period, your performance will be evaluated. The company reserves the right to extend the probation period, if your performance is not satisfactory and not meeting the requisite standards. Employment Terms and conditions and notice period is applicable during the probationary period. "

# P43: Transfer & Deputation
p43 = doc.paragraphs[43]
p43.runs[0].text = "4. Transfer & Deputation:"
p43.runs[1].text = " You will be posted at Harleys India Private Limited at "
p43.runs[2].text = "{{ reporting_location }}"
p43.runs[3].text = " as a "
p43.runs[4].text = "“{{ designation }}”"
p43.runs[5].text = ". However, your services can be transferred to other departments, locations, subsidiaries, sister companies or deputed to any outlet within or outside India based on the requirements of the organization, as decided by the Management. "

# P45: Reporting / Joining
p45 = doc.paragraphs[45]
p45.runs[0].text = "5. Reporting/Joining:"
p45.runs[1].text = " You have to report to Harleys India Private Limited, {{ reporting_location }} at {{ reporting_time }}. In the event of failure of above, this letter of appointment stands withdrawn, unless the Management has agreed to give an extension in writing. "

# P144: Signatory line in Annexure C
p144 = doc.paragraphs[144]
p144.runs[0].text = "{{ signatory_name }}                                                                    Signature of the candidate"

# TABLE 0: Annexure C Compensation Structure
table = doc.tables[0]

# Row 2: Employee Name
# Cells in row 2: [0] Employee Name, [1] Employee Name, [2] Roshan Mokal, [3] Roshan Mokal
for cell in [table.rows[2].cells[2], table.rows[2].cells[3]]:
    for p in cell.paragraphs:
        for r in p.runs:
            r.text = ""
    cell.paragraphs[0].text = "{{ candidate_name }}"

# Row 3: Annual CTC
for cell in [table.rows[3].cells[1], table.rows[3].cells[2], table.rows[3].cells[3]]:
    for p in cell.paragraphs:
        for r in p.runs:
            r.text = ""
    cell.paragraphs[0].text = "{{ salary_amount }}"

# Row 4: Monthly CTC
for cell in [table.rows[4].cells[1], table.rows[4].cells[2], table.rows[4].cells[3]]:
    for p in cell.paragraphs:
        for r in p.runs:
            r.text = ""
    cell.paragraphs[0].text = "₹{{ ctc_monthly }}"

# Row 6: Basic
for p in table.rows[6].cells[2].paragraphs:
    for r in p.runs: r.text = ""
table.rows[6].cells[2].paragraphs[0].text = "₹{{ basic_yearly }}"

for p in table.rows[6].cells[3].paragraphs:
    for r in p.runs: r.text = ""
table.rows[6].cells[3].paragraphs[0].text = "₹{{ basic_monthly }}"

# Row 7: HRA
for p in table.rows[7].cells[2].paragraphs:
    for r in p.runs: r.text = ""
table.rows[7].cells[2].paragraphs[0].text = "₹{{ hra_yearly }}"

for p in table.rows[7].cells[3].paragraphs:
    for r in p.runs: r.text = ""
table.rows[7].cells[3].paragraphs[0].text = "₹{{ hra_monthly }}"

# Row 8: Special Allowance
for p in table.rows[8].cells[2].paragraphs:
    for r in p.runs: r.text = ""
table.rows[8].cells[2].paragraphs[0].text = "₹{{ special_yearly }}"

for p in table.rows[8].cells[3].paragraphs:
    for r in p.runs: r.text = ""
table.rows[8].cells[3].paragraphs[0].text = "₹{{ special_monthly }}"

# Row 9: Total Emoluments
for p in table.rows[9].cells[2].paragraphs:
    for r in p.runs: r.text = ""
table.rows[9].cells[2].paragraphs[0].text = "₹{{ gross_yearly }}"

for p in table.rows[9].cells[3].paragraphs:
    for r in p.runs: r.text = ""
table.rows[9].cells[3].paragraphs[0].text = "₹{{ gross_monthly }}"

# Row 10: Employee PF
for p in table.rows[10].cells[2].paragraphs:
    for r in p.runs: r.text = ""
table.rows[10].cells[2].paragraphs[0].text = "{{ epf_yearly }}"

for p in table.rows[10].cells[3].paragraphs:
    for r in p.runs: r.text = ""
table.rows[10].cells[3].paragraphs[0].text = "{{ epf_monthly }}"

# Row 11: Employer PF
for p in table.rows[11].cells[2].paragraphs:
    for r in p.runs: r.text = ""
table.rows[11].cells[2].paragraphs[0].text = "{{ emp_pf_yearly }}"

for p in table.rows[11].cells[3].paragraphs:
    for r in p.runs: r.text = ""
table.rows[11].cells[3].paragraphs[0].text = "{{ emp_pf_monthly }}"

# Row 12: Professional Tax
for p in table.rows[12].cells[2].paragraphs:
    for r in p.runs: r.text = ""
table.rows[12].cells[2].paragraphs[0].text = "₹{{ pt_yearly }}"

for p in table.rows[12].cells[3].paragraphs:
    for r in p.runs: r.text = ""
table.rows[12].cells[3].paragraphs[0].text = "₹{{ pt_monthly }}"

# Row 13: Total Deductions
for p in table.rows[13].cells[2].paragraphs:
    for r in p.runs: r.text = ""
table.rows[13].cells[2].paragraphs[0].text = "₹{{ total_deductions_yearly }}"

for p in table.rows[13].cells[3].paragraphs:
    for r in p.runs: r.text = ""
table.rows[13].cells[3].paragraphs[0].text = "₹{{ total_deductions_monthly }}"

# Row 14: In Hand Salary
for p in table.rows[14].cells[2].paragraphs:
    for r in p.runs: r.text = ""
table.rows[14].cells[2].paragraphs[0].text = "₹{{ in_hand_yearly }}"

for p in table.rows[14].cells[3].paragraphs:
    for r in p.runs: r.text = ""
table.rows[14].cells[3].paragraphs[0].text = "₹{{ in_hand_monthly }}"

doc.save(template_path)
print(f"Appointment template successfully created and saved to: {template_path}")

