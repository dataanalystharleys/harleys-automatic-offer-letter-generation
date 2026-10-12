import os
import shutil
import docx
from docxtpl import DocxTemplate

source_file = r"C:\Users\harle\Downloads\Offer Letter Harley.docx"
target_dir = r"C:\Users\harle\.gemini\antigravity\scratch\harleys-automatic-offer-letter-generation"
template_path = os.path.join(target_dir, "template_offer_letter.docx")

shutil.copyfile(source_file, template_path)
doc = docx.Document(template_path)

# P0: Date
p0 = doc.paragraphs[0]
p0.runs[0].text = "  Date: {{ offer_date }}"

# P2: Candidate Name
p2 = doc.paragraphs[2]
p2.runs[2].text = "{{ candidate_name }}"

# P3: Address
p3 = doc.paragraphs[3]
p3.runs[0].text = "Address: "
p3.runs[1].text = "{{ address }}"
p3.runs[2].text = ""

# P4: Phone
p4 = doc.paragraphs[4]
p4.runs[0].text = ""
p4.runs[1].text = "{{ phone }}"

# P5: Email
p5 = doc.paragraphs[5]
p5.runs[0].text = "Mail Id: "
p5.runs[1].text = "{{ email }}"
p5.runs[2].text = ""
p5.runs[3].text = ""

# P9: Salutation and Name
p9 = doc.paragraphs[9]
p9.runs[0].text = "Dear {{ salutation }} "
p9.runs[1].text = "{{ candidate_name }}"

# P11: Designation and Company Name
p11 = doc.paragraphs[11]
p11.runs[0].text = "We are pleased to offer you employment in our organization as “"
p11.runs[1].text = "{{ designation }}"
p11.runs[2].text = "” "
p11.runs[3].text = "at “"
p11.runs[4].text = "{{ company_name }}"
p11.runs[5].text = "”."

# P12: Salary
p12 = doc.paragraphs[12]
p12.runs[0].text = "The position carries an annual compensation of "
p12.runs[1].text = "Rs."
p12.runs[2].text = "{{ salary_amount }}"
p12.runs[3].text = ""
p12.runs[4].text = ""
p12.runs[5].text = ""
p12.runs[6].text = ""
p12.runs[7].text = "/- "
p12.runs[8].text = "in Rupees "
p12.runs[9].text = "{{ salary_in_words }} "
p12.runs[10].text = "which includes statutory deductions such as Health card, PT, PFTDS if any."

# P13: Reporting location, time, and joining date
p13 = doc.paragraphs[13]
p13.runs[0].text = "You are requested to report to "
p13.runs[1].text = "{{ reporting_location }}"
p13.runs[2].text = " "
p13.runs[3].text = " at"
p13.runs[4].text = " {{ reporting_time }} on "
p13.runs[5].text = "{{ joining_date }} "
p13.runs[6].text = ""
p13.runs[7].text = ""
p13.runs[8].text = ""
p13.runs[9].text = ""
p13.runs[10].text = "with below "
p13.runs[11].text = "document :"
p13.runs[12].text = "-"

# P24: HR Signatory
p24 = doc.paragraphs[24]
p24.runs[0].text = "{{ hr_signatory }}"

doc.save(template_path)
print(f"Template successfully saved to: {template_path}")

# Now test rendering with docxtpl
tpl = DocxTemplate(template_path)
test_context = {
    "offer_date": "12-Oct-2026",
    "salutation": "Mr.",
    "candidate_name": "Harley Tester",
    "address": "Flat 402, Green Meadows, Hyderabad, Telangana 500081",
    "phone": "+91 9876543210",
    "email": "harley.tester@example.com",
    "designation": "Senior AI Architect",
    "company_name": "HARLEY INDIA PVT LTD",
    "salary_amount": "65,00,000",
    "salary_in_words": "Sixty Five Lakh only/-",
    "reporting_location": "Nanakramguda",
    "reporting_time": "10:00 AM",
    "joining_date": "01-11-2026",
    "hr_signatory": "Sruti Associate - HR"
}

test_output = os.path.join(target_dir, "test_output.docx")
tpl.render(test_context)
tpl.save(test_output)
print(f"Test render successful! Saved to: {test_output}")

# Verify generated document
doc_verify = docx.Document(test_output)
print("\n=== VERIFYING RENDERED OUTPUT ===")
for i, p in enumerate(doc_verify.paragraphs):
    if p.text.strip():
        print(f"P{i}: {p.text}")

